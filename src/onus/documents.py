"""Read the document being checked into blocks: one per text frame,
paragraph, table cell or chart value, each with a location a reader can find.
Pictures are collected alongside, for the model to transcribe (images.py)."""

from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass, field, replace
from pathlib import Path

# A block that is only a short number is a page number or a list marker.
_BARE_NUMBER = re.compile(r"^\s*\d{1,3}\s*$")
_MEDIA_ALIASES = {"image/jpg": "image/jpeg", "image/x-png": "image/png"}
# Alt text a generator fills with the file name says nothing about the picture.
_FILE_NAME = re.compile(r"[\w .()-]+\.(?:png|jpe?g|gif|webp|svg|emf|bmp|tiff?)", re.I)


@dataclass(frozen=True)
class Block:
    id: int
    label: str  # "slide 3", "paragraph 12", "line 40"
    text: str
    origin: str = "text"  # "text" | "chart" | "image": a transcription is only as good as the model's reading
    plain: bool = False  # ends in a chart value printed with no format: a magnitude, no unit


@dataclass(frozen=True)
class Figure:
    label: str  # the slide or paragraph it sits on, as a Block's label
    ordinal: int
    data: bytes
    media_type: str
    alt: str = ""

    @property
    def location(self) -> str:
        return f"{self.label} (image {self.ordinal})"


def _media(content_type: str) -> str:
    return _MEDIA_ALIASES.get(content_type.lower(), content_type.lower())


def _alt(text: str) -> str:
    text = " ".join(text.split())
    if _FILE_NAME.fullmatch(text) or "automatically generated" in text.lower():
        return ""
    return text


@dataclass
class Document:
    path: Path
    format: str
    blocks: list[Block]
    figures: list[Figure] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)

    def extend(self, items: Iterable[Block]) -> None:
        """Add blocks, numbered on from the last, with their text normalised and a repeat dropped."""
        seen = {(b.label, b.text) for b in self.blocks}
        for block in items:
            text = " ".join(block.text.split())
            if not text or _BARE_NUMBER.match(text):
                continue
            if (block.label, text) in seen:  # a slide repeats its eyebrow, a table its header
                continue
            seen.add((block.label, text))
            self.blocks.append(replace(block, id=len(self.blocks), text=text))


def _shapes(shapes):
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _pptx(path: Path) -> tuple[list[Block], list[Figure]]:
    from pptx import Presentation
    from pptx.shapes.picture import Picture

    from onus import charts

    texts: list[Block] = []
    figures: list[Figure] = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        pictures = graphs = 0
        for shape in _shapes(slide.shapes):
            if getattr(shape, "has_chart", False) and shape.has_chart:
                graphs += 1
                where = f"slide {n} (chart {graphs})"
                texts += [Block(0, where, text, "chart", plain) for _, text, plain in charts.blocks(shape.chart, where)]
                continue
            if isinstance(shape, Picture):
                pictures += 1
                # python-pptx's Image wrapper decodes the bytes to name their type; the part already knows it
                rId = shape._element.blip_rId
                if rId is None:  # a linked picture has no bytes in the file
                    continue
                try:
                    part = shape.part.related_part(rId)
                except (KeyError, ValueError):  # a dangling or external relationship
                    continue
                alt = next(iter(shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr")), "")
                figures.append(Figure(f"slide {n}", pictures, part.blob, _media(part.content_type), _alt(alt)))
                continue
            frames = []
            if shape.has_text_frame:
                frames = [shape.text_frame]
            elif getattr(shape, "has_table", False) and shape.has_table:
                frames = [c.text_frame for row in shape.table.rows for c in row.cells]
            for frame in frames:
                for paragraph in frame.text.split("\n"):
                    texts.append(Block(0, f"slide {n}", paragraph))
    return texts, figures


def _docx(path: Path) -> tuple[list[Block], list[Figure]]:
    import docx

    d = docx.Document(str(path))
    texts: list[Block] = []
    figures: list[Figure] = []

    def pictures(label: str, scope) -> None:
        for n, pic in enumerate(scope.xpath(".//pic:pic[not(ancestor::*[local-name()='Fallback'])]"), start=1):
            rids = pic.xpath("./pic:blipFill/a:blip/@r:embed")
            part = d.part.related_parts.get(rids[0]) if rids else None
            if part is None:  # a linked picture has no bytes in the file
                continue
            described = (pic.xpath("./pic:nvPicPr/pic:cNvPr/@descr")
                         or pic.xpath("ancestor::w:drawing[1]//wp:docPr/@descr"))
            alt = _alt(next(iter(described), ""))
            figures.append(Figure(label, n, part.blob, _media(part.content_type), alt))

    for i, p in enumerate(d.paragraphs, start=1):
        texts.append(Block(0, f"paragraph {i}", p.text))
        pictures(f"paragraph {i}", p._p)
    for t, table in enumerate(d.tables, start=1):
        for row in table.rows:
            for cell in row.cells:
                texts.append(Block(0, f"table {t}", cell.text))
        pictures(f"table {t}", table._tbl)  # once per table: a merged cell is listed per column
    return texts, figures


def _text(path: Path) -> tuple[list[Block], list[Figure]]:
    return [Block(0, f"line {i}", line) for i, line in enumerate(path.read_text().splitlines(), start=1)], []


READERS = {".pptx": _pptx, ".docx": _docx, ".md": _text, ".txt": _text, ".markdown": _text}


def load(path: str | Path) -> Document:
    path = Path(path)
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(f"{path.name}: unsupported format (onus reads {', '.join(sorted(READERS))})")
    texts, figures = reader(path)
    doc = Document(path, path.suffix.lower().lstrip("."), [], figures)
    doc.extend(texts)
    return doc
