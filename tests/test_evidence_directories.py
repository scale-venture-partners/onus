"""A directory -- a git repository, usually -- as evidence: what is read, how passages
are ranked for the judge, and when a number in code may support a claim."""

import json
import subprocess

import pytest

from onus import claims, cli, documents, verify
from onus import evidence as ev


def repo(tmp_path, files, git=True):
    root = tmp_path / "proj"
    for rel, data in files.items():
        path = root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data if isinstance(data, bytes) else data.encode())
    if git:
        run = lambda *a: subprocess.run(["git", "-C", str(root), *a], check=True, capture_output=True)  # noqa: E731
        run("init", "-q")
        run("add", "-A")
    return root


FILES = {
    "README.md": "The harness stops after 40 tool-calling turns.\n\nThree linters run as hooks.",
    "loop.py": "MAX_TURNS = 40\nPORT = 5055\n",
    "docs/notes.txt": "Workshop on 2026-10-02.",
    "data/metrics.csv": "month,arr\nsep,2.4\n",
    "uv.lock": "version = 1\nrevision = 99\n",
    "logo.png": b"\x89PNG\r\n\x1a\n\x00\x00binary",
    "latin1.txt": "caf\xe9".encode("latin-1"),
    "tool.bin": "not a known type",
    "app.min.js": "var a=1",
    "LICENSE": "MIT License",
}


def sources(e):
    return sorted({p.source.split("#")[0] for p in e.passages})


def test_a_repository_contributes_its_tracked_text_files_by_relative_path(tmp_path):
    root = repo(tmp_path, {**FILES, ".gitignore": "out/\n", "out/deck_notes.md": "Generated: 99 slides"})
    e = ev.load([root])
    # out/ is gitignored, so never read; nor are the lockfile, the image, the binary and the
    # minified bundle. Not valid UTF-8 counts as binary.
    assert sources(e) == ["proj/LICENSE", "proj/README.md", "proj/data/metrics.csv", "proj/docs/notes.txt",
                          "proj/loop.py"]
    assert e.notes == ["evidence: 5 tracked files from proj/ (skipped 4 other types, 1 binary, 1 lockfiles)"]
    assert {p.source: p.kind for p in e.passages}["proj/loop.py"] == "code"


def test_outside_git_a_directory_is_walked_without_tool_folders(tmp_path):
    root = repo(tmp_path, {**FILES, ".venv/lib/x.py": "X = 1", "node_modules/a/index.js": "1",
                           "__pycache__/m.txt": "cache", ".hidden/n.md": "hidden"}, git=False)
    e = ev.load([root])
    assert "proj/README.md" in sources(e)
    assert not any(s.startswith(("proj/.venv", "proj/node_modules", "proj/__pycache__", "proj/.hidden"))
                   for s in sources(e))
    assert e.notes[0].startswith("evidence: 5 files from proj/")


def test_excludes_large_files_and_documents(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches

    root = repo(tmp_path, {**FILES, "tests/test_x.py": "assert 40", "big.md": "x " * 200}, git=False)
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "NRR is 118%"
    prs.save(root / "deck.pptx")
    monkeypatch.setattr(ev, "MAX_FILE_BYTES", 300)
    e = ev.load([root], exclude=["tests", "*.csv"])
    assert "proj/tests/test_x.py" not in sources(e) and "proj/data/metrics.csv" not in sources(e)
    assert "proj/big.md" not in sources(e)
    assert any(p.source == "proj/deck.pptx" and "118%" in p.text for p in e.passages)
    assert "2 excluded" in e.notes[0] and "1 too large" in e.notes[0]


def test_one_malformed_file_is_skipped_with_a_note(tmp_path):
    root = repo(tmp_path, {"README.md": "ok", "bad.json": "{nope"}, git=False)
    e = ev.load([root])
    assert sources(e) == ["proj/README.md"]
    assert e.notes[0].startswith("evidence: skipped proj/bad.json (JSONDecodeError")
    assert e.notes[1] == "evidence: 1 files from proj/ (skipped 1 unreadable)"


def test_a_transcript_and_json_inside_a_directory_are_read_as_such(tmp_path):
    lines = [{"role": "user", "content": "ARR $2.4B"}, {"role": "assistant", "content": "ARR $9B"}]
    root = repo(tmp_path, {"s.jsonl": "\n".join(json.dumps(x) for x in lines),
                           "d.json": json.dumps({"nrr": "118%"})}, git=False)
    text = " ".join(p.text for p in ev.load([root]).passages)
    assert "$2.4B" in text and "$9B" not in text and "118%" in text


def test_without_git_installed_the_walk_is_used(tmp_path, monkeypatch):
    root = repo(tmp_path, {"README.md": "ok"}, git=False)

    def no_git(*a, **kw):
        raise FileNotFoundError("git")
    monkeypatch.setattr(ev.subprocess, "run", no_git)
    assert sources(ev.load([root])) == ["proj/README.md"]


# -- ranking ----------------------------------------------------------------------------------

def test_prose_outranks_code_that_matches_as_well():
    e = ev.Evidence()
    e.add("loop.py", "def stop_after_turns(turns): return turns > MAX_TURNS", "code")
    e.add("README.md", "The loop stops after a fixed number of turns.")
    assert e.relevant("loop stops after turns", k=1)[0].source == "README.md"


def test_a_rare_word_outweighs_common_ones():
    e = ev.Evidence()
    for i in range(20):
        e.add(f"f{i}.md", f"The model runs tests on the deck, take {i}.")
    e.add("overset.md", "overset measures the render.")
    assert e.relevant("the model measures overset tests", k=1)[0].source == "overset.md"


def test_identifiers_are_split_into_words():
    assert ev._tokens("MAX_TURNS maxTurns") == ["max", "turns", "max", "turns"]


def test_the_index_is_rebuilt_when_evidence_grows():
    e = ev.Evidence()
    e.add("a.md", "alpha")
    assert e.relevant("beta", k=1)[0].source == "a.md", "nothing matches: the first passages, as before"
    e.add("b.md", "beta")
    assert e.relevant("beta", k=1)[0].source == "b.md"


# -- numbers in code ----------------------------------------------------------------------------

def claim_of(tmp_path, text):
    d = documents.load((tmp_path / "c.md").write_text(text) and tmp_path / "c.md")
    return claims.scan_claims(d)[0]


@pytest.mark.parametrize(("claim", "status"), [
    ("The loop allows 40 turns.", "supported"),       # MAX_TURNS = 40 names turns
    ("We have 5055 customers.", "unsupported"),       # PORT = 5055 is not about customers
])
def test_a_number_in_code_supports_a_claim_only_about_the_same_thing(tmp_path, claim, status):
    e = ev.Evidence()
    e.add("loop.py", "MAX_TURNS = 40\nPORT = 5055", "code")
    c = claim_of(tmp_path, claim)
    verify.check_numbers(c, e)
    assert c.status == status


@pytest.mark.parametrize(("claim", "status"), [
    ("Each deck took 113 tool calls.", "unsupported"),  # "113 unsupported claims" is about claims
    ("The deck had 113 unsupported claims.", "supported"),
    ("A run costs $17.", "unsupported"),                # "on slide 17"
])
def test_a_number_read_from_a_directory_needs_its_sentence_to_be_about_the_claim(tmp_path, claim, status):
    root = repo(tmp_path, {"plan.md": "The first check found 113 unsupported claims. Riff caught the order "
                                      "(on slide 17).\n\nRun costs vary."}, git=False)
    c = claim_of(tmp_path, claim)
    verify.check_numbers(c, ev.load([root]))
    assert c.status == status


def test_a_file_named_on_its_own_keeps_the_looser_rule(tmp_path):
    (tmp_path / "plan.md").write_text("There were 17 of them.")
    c = claim_of(tmp_path, "We shipped 17 decks.")
    verify.check_numbers(c, ev.load([tmp_path / "plan.md"]))
    assert c.status == "supported", "the user chose this file: any matching number counts, as before"


def test_a_table_row_counts_its_header_as_context(tmp_path):
    root = repo(tmp_path, {"metrics.tsv": "month\tcustomers\tchurn\nsep\t240\t3\n"}, git=False)
    e = ev.load([root])
    for text, status in (("We have 240 customers.", "supported"), ("We ran 240 tests.", "unsupported")):
        c = claim_of(tmp_path, text)
        verify.check_numbers(c, e)
        assert c.status == status, text


def test_code_licenses_no_arithmetic(tmp_path):
    e = ev.Evidence()
    e.add("loop.py", "turns went from 40 to 80 here", "code")
    assert e.passages[0].derived == []
    e.add("notes.md", "turns went from 40 to 80 here")
    assert e.passages[1].derived


# -- the command line -----------------------------------------------------------------------------

def test_the_cli_takes_a_directory_and_excludes(tmp_path, capsys):
    root = repo(tmp_path, FILES)
    (tmp_path / "deck.md").write_text("The harness stops after 40 tool-calling turns.\n")
    assert cli.main([str(tmp_path / "deck.md"), "-e", str(root), "--evidence-exclude", "docs/*", "--no-llm"]) == 0
    out = capsys.readouterr().out
    assert "evidence: 4 tracked files from proj/ (skipped 3 other types, 1 binary, 1 excluded, 1 lockfiles)" in out


def test_missing_evidence_is_a_usage_error(tmp_path, capsys):
    (tmp_path / "deck.md").write_text("ARR $2.4B\n")
    assert cli.main([str(tmp_path / "deck.md"), "-e", str(tmp_path / "gone"), "--no-llm"]) == 2
    assert "evidence" in capsys.readouterr().err
