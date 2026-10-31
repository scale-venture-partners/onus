"""Native charts: read deterministically, one block per value, printed as PowerPoint would."""

import pytest

from onus import charts, claims, documents, numbers
from onus import evidence as ev
from onus.documents import Block, Document
from onus.engine import check
from onus.settings import Settings


def chart_deck(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = (Inches(1), Inches(1), Inches(6), Inches(4))
    data = CategoryChartData()
    data.categories = ["Before", "Now", "2026 target"]
    data.add_series("Backlog", (2.1, 1.4, 1.2), number_format='0.0"x"')
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *box, data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Repair backlog"
    chart.plots[0].series[0].points[2].data_label.text_frame.text = "Target 1.2x"
    data = CategoryChartData()
    data.categories = ["Fare revenue", "Renewal"]
    data.add_series("This year", (2400000, 0.88))
    s.shapes.add_group_shape().shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, *box, data)
    xy = XyChartData()
    xy.add_series("Docks").add_data_point(2, 24)
    s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, *box, xy)
    prs.save(tmp_path / "c.pptx")
    return tmp_path / "c.pptx"


def test_native_charts_are_read_as_category_value_pairs(tmp_path):
    d = documents.load(chart_deck(tmp_path))
    assert [(b.label, b.text, b.plain) for b in d.blocks] == [
        ("slide 1 (chart 1)", "Repair backlog", False),
        ("slide 1 (chart 1)", "Backlog, Before: 2.1x", False),
        ("slide 1 (chart 1)", "Backlog, Now: 1.4x", False),
        ("slide 1 (chart 1)", "Backlog, 2026 target: 1.2x", False),
        ("slide 1 (chart 1)", "Target 1.2x", False),
        ("slide 1 (chart 2)", "This year, Fare revenue: 2400000", True),
        ("slide 1 (chart 2)", "This year, Renewal: 0.88", True),
        ("slide 1 (chart 3)", "Docks: 24", True),
    ], "a chart in a group is read; a value with no format is marked as a bare magnitude"


def test_a_deck_with_charts_is_evidence_one_passage_per_chart(tmp_path):
    e = ev.load([chart_deck(tmp_path)])
    assert [p.source for p in e.passages] == [f"c.pptx:slide 1 (chart {n})" for n in (1, 2, 3)]
    memo = tmp_path / "memo.md"
    memo.write_text("Fare revenue reached $2.4M.\n\nRenewal is 88%.\n\nBacklog is 1.4x.\n\nRenewal was 60%.\n")
    result = check(memo, e, Settings(llm=False))
    assert [(f.code, f.message) for f in result.findings] == [("ONS001", "60%: no number like it in the evidence")]
    supported = [c for c in result.claims if c.status == "supported"]
    assert [c.evidence_source for c in supported] == [f"c.pptx:slide 1 (chart {n})" for n in (2, 2, 1)]


def test_chart_values_are_checked_by_tier_1_with_no_model(tmp_path):
    e = ev.load(texts=["Fare revenue $2.4M; renewal 88%; backlog 1.4x, down from 2.1x; 24 docks; plans for 2026."])
    result = check(chart_deck(tmp_path), e, Settings(llm=False))
    assert [(f.code, f.label, f.snippet) for f in result.findings] == [
        ("ONS001", "slide 1 (chart 1)", "Backlog, 2026 target: 1.2x"), ("ONS001", "slide 1 (chart 1)", "Target 1.2x")]
    assert not result.used_model


def test_the_format_comes_from_data_labels_then_the_series_then_the_axis(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = (Inches(1), Inches(1), Inches(6), Inches(4))
    data = CategoryChartData()
    data.categories = ["Renewed", "Lapsed"]
    data.add_series("Members", (0.88, None))
    s.shapes.add_chart(XL_CHART_TYPE.PIE, *box, data)
    labelled = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *box, data).chart
    labelled.plots[0].has_data_labels = True
    labelled.plots[0].data_labels.number_format = "0%"
    labelled.plots[0].data_labels.number_format_is_linked = False
    axis = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *box, data).chart
    axis.value_axis.tick_labels.number_format = "0.0%"
    axis.value_axis.tick_labels.number_format_is_linked = False
    prs.save(tmp_path / "f.pptx")
    assert [(b.text, b.plain) for b in documents.load(tmp_path / "f.pptx").blocks] == [
        ("Members, Renewed: 0.88", True), ("Members, Renewed: 88%", False), ("Members, Renewed: 88.0%", False)], \
        "a pie has no value axis to fall back on; an empty point is not a value"


def test_dates_and_numbers_on_the_category_axis_are_written_as_printed(tmp_path):
    from datetime import date

    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = (Inches(1), Inches(1), Inches(6), Inches(4))
    data = CategoryChartData()
    data.categories = [date(2025, 1, 1), date(2025, 4, 1)]
    data.add_series("Fare revenue", (2.4, 2.6), number_format='"$"0.0"M"')
    s.shapes.add_chart(XL_CHART_TYPE.LINE, *box, data)
    data = CategoryChartData()
    data.categories = [date(2025, 7, 1)]
    data.categories.number_format = "mmm yyyy"
    data.add_series("Fare revenue", (2.9,), number_format='"$"0.0"M"')
    s.shapes.add_chart(XL_CHART_TYPE.LINE, *box, data)
    data = CategoryChartData()
    data.categories = [2023, 2024]
    data.add_series("Margin", (-0.05, 0.12), number_format="0%")
    labelled = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *box, data).chart
    labelled.plots[0].series[0].data_labels.number_format = "0.0%"
    labelled.plots[0].series[0].data_labels.number_format_is_linked = False
    prs.save(tmp_path / "d.pptx")
    assert [b.text for b in documents.load(tmp_path / "d.pptx").blocks] == [
        "Fare revenue, 1 January 2025: $2.4M", "Fare revenue, 1 April 2025: $2.6M", "Fare revenue, July 2025: $2.9M",
        "Margin, 2023: minus 5.0%", "Margin, 2024: 12.0%"], \
        "a date serial is never a count; a series' own label format wins; a negative keeps a sign the scanner reads"


@pytest.mark.parametrize("value, fmt, printed", [
    (2.1, '0.0"x"', "2.1x"),
    (2400000, '"$"#,##0.0,,"M"', "$2.4M"),
    (2400000, "$#,##0", "$2,400,000"),
    (1200, '#,##0.0,"K"', "1.2K"),
    (0.88, "0%", "88%"),
    (0.884, "0.0%", "88.4%"),
    (24, '0 "docks"', "24 docks"),
    (1234.5, "[$€-2] #,##0.00", "€ 1,234.50"),
    (-1.5, "0.0;-0.0", "minus 1.5"),
    (-0.05, "0%", "minus 5%"),
    (15000000000, None, "15000000000"),
    (0.00005, None, "0.00005"),
    (0.1 + 0.2, None, "0.3"),
    (2.1, '0.0"x', "2.1"),
    (2.1, "0.0\\x", "2.1x"),
    (2.1, "[Red]_(0.0_)", "2.1"),
    (2400000, "General", "2400000"),
    (0.88, None, "0.88"),
])
def test_number_formats_render_as_powerpoint_prints_them(value, fmt, printed):
    assert charts.render(value, fmt) == printed


def test_a_chart_passage_supports_a_claim_only_from_the_value_s_own_line():
    passage = ev.Passage("deck.pptx:slide 1 (chart 1)", "Docks: 24\nBudget: 17", strict=True, plain=True)
    seventeen = next(q for q in passage.quantities if q.value == 17)
    assert not passage.can_support(seventeen, "we run 17 docks")
    assert passage.can_support(seventeen, "a budget of 17")


def test_a_model_span_cut_short_of_the_value_keeps_the_category_s_year(tmp_path):
    d = Document(tmp_path / "d.pptx", "pptx", [Block(0, "slide 1 (chart 1)", "Stations, 2024: 1200", "chart", True)])
    extracted = claims.ExtractedClaim(block=0, text="Stations, 2024", kind="temporal", subject="")
    (short, rest), (missed,), _ = claims.merge(d, [extracted])
    assert [(q.raw, q.kind) for q in short.quantities] == [("2024", "year")]
    assert rest is missed and [(q.raw, q.kind) for q in missed.quantities] == [("1200", "number")]
    (whole,), _, _ = claims.merge(d, [claims.ExtractedClaim(block=0, text="Stations, 2024: 1200", kind="quantitative",
                                                               subject="")])
    assert [(q.raw, q.kind) for q in whole.quantities] == [("2024", "year"), ("1200", "number")]
    assert numbers.find("Margin: minus 1.5", plain=True)[0].value == 1.5
