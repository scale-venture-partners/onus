"""Evidence, claim extraction and the two verification tiers, with the model
replaced by pydantic-ai's TestModel and FunctionModel."""

import base64
import io
import json
import zipfile
from pathlib import Path

from pydantic_ai import BinaryContent
from pydantic_ai.messages import ModelResponse, ToolCallPart
from pydantic_ai.models.function import AgentInfo, FunctionModel
from pydantic_ai.models.test import TestModel

from onus import claims, documents, images, verify
from onus import evidence as ev
from onus.documents import Block, Document, Figure
from onus.engine import Finding, Result, check
from onus.settings import Settings

FACTS = ("Fare revenue $2.4B, up 42% year over year; median repair backlog 1.4x, down from 2.1x; "
         "24 active stations. Plans: east bank, night shuttle, and cargo bikes.")
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5Erk"
                       "Jggg==")


def doc(tmp_path, text, name="memo.md"):
    path = tmp_path / name
    path.write_text(text)
    return documents.load(path)


def facts():
    e = ev.Evidence()
    e.add("facts.md", FACTS)
    return e


# -- documents ----------------------------------------------------------------

def test_page_numbers_and_markers_are_not_blocks(tmp_path):
    d = doc(tmp_path, "Revenue grew\n\n3\n01\nRevenue grew\n")
    assert [(b.label, b.text) for b in d.blocks] == [("line 1", "Revenue grew"), ("line 5", "Revenue grew")], \
        "two identical lines are two places; only bare numbers go"


def test_decks_are_read_per_slide(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for text in ("ARR is $2.4B", "NRR is 118%"):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = text
    prs.save(tmp_path / "d.pptx")
    d = documents.load(tmp_path / "d.pptx")
    assert [(b.label, b.text) for b in d.blocks] == [("slide 1", "ARR is $2.4B"), ("slide 2", "NRR is 118%")]


def test_grouped_shapes_are_read_through_their_groups(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = (Inches(1), Inches(1), Inches(4), Inches(1))
    tile = s.shapes.add_group_shape()
    tile.shapes.add_textbox(*box).text_frame.text = "ARR reached $9.7M"
    tile.shapes.add_group_shape().shapes.add_textbox(*box).text_frame.text = "Payback is 14 months"
    s.shapes.add_textbox(*box).text_frame.text = "Churn fell to 3%"
    prs.save(tmp_path / "d.pptx")
    d = documents.load(tmp_path / "d.pptx")
    assert [b.text for b in d.blocks] == ["ARR reached $9.7M", "Payback is 14 months", "Churn fell to 3%"]
    assert "$9.7M" in ev.load([tmp_path / "d.pptx"]).passages[0].text


def deck(tmp_path, *shapes):
    """A one-slide deck: a text box per string, a picture per bytes."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in shapes:
        if isinstance(shape, str):
            s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = shape
        else:
            s.shapes.add_picture(io.BytesIO(shape), Inches(1), Inches(2), width=Inches(5), height=Inches(3))
    prs.save(tmp_path / "d.pptx")
    return tmp_path / "d.pptx"


def test_pictures_are_collected_with_their_place(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = deck(tmp_path, "Network results", PNG)
    prs = Presentation(path)
    prs.slides[0].shapes.add_group_shape().shapes.add_picture(io.BytesIO(PNG + b"\0"), Inches(7), Inches(1), Inches(1))
    prs.save(path)
    d = documents.load(path)
    assert [b.text for b in d.blocks] == ["Network results"]
    assert [(f.location, f.media_type, f.data) for f in d.figures] == [
        ("slide 1 (image 1)", "image/png", PNG), ("slide 1 (image 2)", "image/png", PNG + b"\0")]

    import docx

    w = docx.Document()
    w.add_paragraph("ARR reached $2.4B.")
    w.add_picture(io.BytesIO(PNG), width=Inches(4), height=Inches(2))
    table = w.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.paragraphs[0].add_run().add_picture(io.BytesIO(PNG), width=Inches(2), height=Inches(1))
    w.save(tmp_path / "m.docx")
    d = documents.load(tmp_path / "m.docx")
    assert [(f.location, f.media_type, f.data, f.alt) for f in d.figures] == [
        ("paragraph 2 (image 1)", "image/png", PNG, ""), ("table 1 (image 1)", "image/png", PNG, "")], \
        "a picture in a merged cell is read once; a file-name alt text is no alt text"


def repack(path, media, blob, content_type):
    """Swap a deck's picture for bytes python-pptx cannot decode, declared as PowerPoint would."""
    with zipfile.ZipFile(path) as z:
        items = {name: z.read(name) for name in z.namelist()}
    items[media] = blob
    override = f'<Override PartName="/{media}" ContentType="{content_type}"/></Types>'.encode()
    items["[Content_Types].xml"] = items["[Content_Types].xml"].replace(b"</Types>", override)
    with zipfile.ZipFile(path, "w") as z:
        for name, data in items.items():
            z.writestr(name, data)


def unlink(path, index=0):
    """Make a saved deck's picture a link to a file outside it, as a deck pasted from the web may be."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation(path)
    blip = [s for s in prs.slides[0].shapes if s.shape_type == 13][index]._element.blipFill.blip
    del blip.attrib[qn("r:embed")]
    blip.set(qn("r:link"), "rId99")
    prs.save(path)


def test_picture_types_are_read_as_declared_not_sniffed(tmp_path):
    path = deck(tmp_path, PNG, PNG + b"\0", PNG + b"\0\0", PNG + b"\0\0\0")
    unlink(path, 1)
    repack(path, "ppt/media/image1.png", b"RIFF....WEBPVP8 ", "image/webp")
    repack(path, "ppt/media/image3.png", b"\1\0\0\0", "image/x-emf")
    repack(path, "ppt/media/image4.png", PNG, "image/JPG")
    d = documents.load(path)
    assert [(f.location, f.media_type, f.data) for f in d.figures] == [
        ("slide 1 (image 1)", "image/webp", b"RIFF....WEBPVP8 "), ("slide 1 (image 3)", "image/x-emf", b"\1\0\0\0"),
        ("slide 1 (image 4)", "image/jpeg", PNG)], \
        "a linked picture is skipped but counted; WebP, EMF and a generator's image/JPG are read as declared"
    assert ev.load([path]).passages == []


def test_pictures_are_read_without_decoding_them(tmp_path, monkeypatch):
    import docx
    from pptx.util import Inches

    path = deck(tmp_path, "Network results", PNG)
    w = docx.Document()
    w.add_picture(io.BytesIO(PNG), width=Inches(4), height=Inches(2))
    w.save(tmp_path / "m.docx")

    def decoded(*a, **kw):
        raise AssertionError("decoded a picture")
    monkeypatch.setattr("PIL.Image.open", decoded)
    assert [(f.location, f.media_type, f.data) for f in documents.load(path).figures] == [
        ("slide 1 (image 1)", "image/png", PNG)]
    assert [(f.location, f.media_type, f.data) for f in documents.load(tmp_path / "m.docx").figures] == [
        ("paragraph 1 (image 1)", "image/png", PNG)]


def test_word_pictures_in_text_boxes_and_groups_are_read_once_each(tmp_path):
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from pptx.util import Inches

    def wrap(paragraph, uri, holder, alt):
        """Move the paragraph's pictures into one drawing of another kind, as Word nests them."""
        drawings = paragraph._p.xpath(".//w:drawing")
        outer = parse_xml(f'<w:drawing {nsdecls("w", "wp", "a")}><wp:inline><wp:docPr id="9" name="x" descr="{alt}"/>'
                          f'<a:graphic><a:graphicData uri="{uri}">{holder}</a:graphicData></a:graphic>'
                          '</wp:inline></w:drawing>')
        inner = outer.xpath(".//*[local-name()='txbxContent' or local-name()='wgp']")[0]
        for drawing in drawings:
            inner.append(drawing if "wgp" in holder else parse_xml(f'<w:p {nsdecls("w")}><w:r/></w:p>'))
            if "wgp" not in holder:
                inner[-1][0].append(drawing)
            else:
                inner[-1].getparent().replace(inner[-1], drawing.xpath(".//pic:pic")[0])
        paragraph._p.xpath("./w:r")[0].append(outer)

    w = docx.Document()
    boxed = w.add_paragraph()
    boxed.add_run().add_picture(io.BytesIO(PNG), width=Inches(3))
    boxed._p.xpath(".//wp:docPr")[0].set("descr", "a chart")
    wrap(boxed, "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
         '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:txbx>'
         '<w:txbxContent/></wps:txbx></wps:wsp>', "a text box")
    grouped = w.add_paragraph()
    grouped.add_run().add_picture(io.BytesIO(PNG + b"\0"), width=Inches(1))
    grouped.add_run().add_picture(io.BytesIO(PNG + b"\0\0"), width=Inches(1))
    wrap(grouped, "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
         '<wpg:wgp xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"/>', "two logos")
    w.save(tmp_path / "n.docx")
    assert [(f.location, f.data, f.alt) for f in documents.load(tmp_path / "n.docx").figures] == [
        ("paragraph 1 (image 1)", PNG, "a chart"),
        ("paragraph 2 (image 1)", PNG + b"\0", "two logos"), ("paragraph 2 (image 2)", PNG + b"\0\0", "two logos")]


def test_a_word_text_box_with_a_fallback_branch_is_read_once(tmp_path):
    import copy

    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from pptx.util import Inches

    w = docx.Document()
    p = w.add_paragraph()
    p.add_run().add_picture(io.BytesIO(PNG), width=Inches(3))
    drawing = p._p.xpath(".//w:drawing")[0]
    holder = parse_xml(
        f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" {nsdecls("w")}'
        ' xmlns:v="urn:schemas-microsoft-com:vml"><mc:Choice Requires="wps"><w:txbxContent/></mc:Choice>'
        '<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent/></v:textbox></v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent>')
    for content in holder.iter(qn("w:txbxContent")):
        content.append(parse_xml(f'<w:p {nsdecls("w")}><w:r/></w:p>'))
        content[0][0].append(copy.deepcopy(drawing))
    drawing.getparent().replace(drawing, holder)
    w.save(tmp_path / "b.docx")
    figures = documents.load(tmp_path / "b.docx").figures
    assert [(f.location, f.data) for f in figures] == [("paragraph 1 (image 1)", PNG)]


def test_alt_text_that_says_nothing_is_dropped():
    assert documents._alt("Screenshot (1).PNG") == ""
    assert documents._alt("A picture containing text, screenshot\n\nDescription automatically generated") == ""
    assert documents._alt("  Repair backlog\nby quarter ") == "Repair backlog by quarter"


def test_a_linked_picture_is_skipped(tmp_path):
    path = deck(tmp_path, "Network results", PNG)
    unlink(path)
    d = documents.load(path)
    assert [b.text for b in d.blocks] == ["Network results"] and d.figures == []
    assert ev.load([path]).passages[0].text == "Network results"


def test_an_unknown_format_is_an_error(tmp_path):
    (tmp_path / "x.pdf").write_text("")
    try:
        documents.load(tmp_path / "x.pdf")
    except ValueError as e:
        assert "unsupported format" in str(e)
    else:
        raise AssertionError


# -- images -------------------------------------------------------------------

def figure(label, data=PNG, media="image/png", alt=""):
    return Figure(label, 1, data, media, alt)


def transcribed(tool, kind, *lines):
    return ModelResponse(parts=[ToolCallPart(tool, {"kind": kind, "lines": list(lines)})])


def test_pictures_are_transcribed_into_blocks_under_their_location(tmp_path):
    calls = []

    def answer(messages, info: AgentInfo):
        calls.append(messages[0].parts[-1].content)
        return transcribed(info.output_tools[0].name, "chart", "Network  results", "Fare revenue, 2025: $3.1M",
                           "Renewal: 88%", "Renewal: 88%")

    d = Document(tmp_path / "d.pptx", "pptx", [Block(0, "slide 1", "Network results")],
                 [figure("slide 1", alt="Repair backlog chart"), figure("slide 2")])
    notes = images.transcribe(d, FunctionModel(answer))
    assert [(b.label, b.text) for b in d.blocks[1:]] == [
        ("slide 1 (image 1)", "Fare revenue, 2025: $3.1M"), ("slide 1 (image 1)", "Renewal: 88%"),
        ("slide 2 (image 1)", "Network results"),
        ("slide 2 (image 1)", "Fare revenue, 2025: $3.1M"), ("slide 2 (image 1)", "Renewal: 88%")], \
        "a line the slide already prints as text is not added again under its picture"
    assert len(calls) == 1, "a picture repeated across slides is read once"
    image, text = calls[0]
    assert isinstance(image, BinaryContent) and image.data == PNG and image.media_type == "image/png"
    assert text.startswith("slide 1 (image 1) of d.pptx") and "Alt text: Repair backlog chart\nNetwork results" in text
    assert notes == ["images: transcribed 2 of 2"]


def test_pictures_the_model_cannot_read_are_skipped_with_a_note(tmp_path):
    seen = []

    def answer(messages, info: AgentInfo):
        seen.append(messages[0].parts[-1].content[0].data)
        if seen[-1] == PNG:
            raise RuntimeError("boom")
        return transcribed(info.output_tools[0].name, "none")

    d = Document(tmp_path / "d.pptx", "pptx", [], [
        figure("slide 1"), figure("slide 2", data=PNG + b"\0"), figure("slide 3", media="image/x-emf"),
        figure("slide 4", data=b"x" * (images.MAX_IMAGE_BYTES + 1))])
    notes = images.transcribe(d, FunctionModel(answer))
    assert d.blocks == [] and len(seen) == 2
    assert notes == ["images: slide 1 (image 1) not read (RuntimeError: boom)",
                     "images: transcribed 0 of 4 (1 image/x-emf, 1 nothing to check, 1 over the size limit)"]


def test_a_chart_image_is_checked_like_the_text_around_it(tmp_path):
    from onus.report import as_dict

    path = deck(tmp_path, "Fare revenue $2.4M", PNG)

    def answer(messages, info: AgentInfo):
        tool, content = info.output_tools[0].name, messages[0].parts[-1].content
        if isinstance(content, list):
            return transcribed(tool, "chart", "Fare revenue, 2026: $3.1M")
        return ModelResponse(parts=[ToolCallPart(tool, {"claims": []})])

    e = ev.load(texts=["Fare revenue $2.4M in 2026."])
    result = check(path, e, Settings(), model=FunctionModel(answer))
    assert [(f.code, f.label, f.snippet) for f in result.findings] == [
        ("ONS001", "slide 1 (image 1)", "Fare revenue, 2026: $3.1M")]
    assert result.notes == ["images: transcribed 1 of 1"]
    assert [(c["origin"], c["status"]) for c in as_dict(result)["claims"]] == [
        ("text", "supported"), ("image", "unsupported")]
    for settings, why in ((Settings(llm=False), "--no-llm"), (Settings(images=False), "--no-images")):
        result = check(path, e, settings, model=FunctionModel(answer))
        assert result.findings == [] and result.notes == [f"images: 1 not read ({why})"]


def test_findings_sort_by_slide_then_picture():
    labels = ["slide 3 (image 2)", "slide 3", "slide 3 (image 1)", "slide 2"]
    r = Result(Path("d.pptx"), [], [Finding("ONS001", "", label, "", "error") for label in labels], 0, False)
    assert [f.label for f in r.sorted()] == ["slide 2", "slide 3", "slide 3 (image 1)", "slide 3 (image 2)"]


# -- evidence -----------------------------------------------------------------

def transcript(tmp_path, lines):
    path = tmp_path / "session.jsonl"
    path.write_text("\n".join(json.dumps(x) for x in lines))
    return path


def test_a_transcript_counts_what_came_in_not_what_the_agent_said(tmp_path):
    path = transcript(tmp_path, [
        {"role": "user", "content": "Facts: ARR $2.4B."},
        {"role": "assistant", "content": [
            {"type": "text", "text": "Eight of 24 stations sell into regulated industries."},
            {"type": "tool_use", "id": "a", "name": "search_crm", "input": {}},
            {"type": "tool_use", "id": "b", "name": "read", "input": {"path": "out/build_deck.py"}},
        ]},
        {"role": "user", "content": [
            {"type": "tool_result", "tool_use_id": "a", "content": "NRR 118% across the network"},
            {"type": "tool_result", "tool_use_id": "b", "content": "deck.cover(title='Eight of 24...')"},
            {"type": "text", "text": "<system-reminder>riff: slide 3 'Eight of our 24'</system-reminder>"},
        ]},
    ])
    e = ev.load([path])
    text = " ".join(p.text for p in e.passages)
    assert "$2.4B" in text and "118%" in text
    assert "Eight" not in text, "the agent's words, its own files read back, and hook reminders are not evidence"
    assert {p.source.split(":")[1] for p in e.passages} == {"user", "search_crm"}


def test_excluded_tools_are_configurable(tmp_path):
    path = transcript(tmp_path, [
        {"role": "assistant", "content": [{"type": "tool_use", "id": "a", "name": "read", "input": {}}]},
        {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "a", "content": "dataroom: 118%"}]},
    ])
    assert ev.load([path]).passages == []
    assert "118%" in ev.load([path], exclude_tools=()).passages[0].text


def test_json_and_other_files_are_flattened_to_text(tmp_path):
    (tmp_path / "d.json").write_text(json.dumps({"arr": "$2.4B", "rows": [{"nrr": 118}, None, True]}))
    (tmp_path / "g.jsonl").write_text('{"note": "backlog 1.4x"}\n')
    e = ev.load([tmp_path / "d.json", tmp_path / "g.jsonl"], texts=["24 stations"])
    text = " ".join(p.text for p in e.passages)
    assert "$2.4B" in text and "118" in text and "backlog 1.4x" in text and "24 stations" in text


def test_long_evidence_is_chunked_on_paragraphs():
    e = ev.Evidence()
    e.add("big.md", "\n\n".join(f"Paragraph {i} " + "x" * 400 for i in range(10)))
    assert len(e.passages) > 1 and e.passages[1].source == "big.md#2"
    assert all(p.text.startswith("Paragraph") for p in e.passages)


def test_a_quote_is_found_verbatim_or_not_at_all():
    e = facts()
    assert e.contains("median repair backlog 1.4x,   down from 2.1x") is not None
    assert e.contains("repair backlog of 1.4") is None
    assert e.contains("") is None


def test_a_quote_ignores_markup_and_punctuation_but_not_words_or_numbers():
    e = ev.load(texts=["**Run 4 ran out of turns on a bug in the skill.** the linter kept reporting the "
                             "`c.pill()` overflow. The audience is the city's planners, in a workshop. "
                             "ARR grew to $2.4B, up 38%."])
    assert e.contains("Run 4 ran out of turns on a bug in the skill. the linter kept reporting the c.pill() overflow")
    assert e.contains("The audience is the city\u2019s planners in a workshop")
    assert e.contains("ARR grew to $2.4B -- up 38%")
    assert e.contains("ARR grew to $2.5B") is None, "a number must match exactly"
    assert e.contains("ARR grew to $2.4B, up 3%") is None
    assert e.contains("Run 5 ran out of turns") is None
    assert e.contains("ran out of turns on a bug in the skills") is None, "whole words only"
    assert e.contains("in the skill the linter kept reporting"), "across a markup boundary"
    assert e.contains("... ** ``") is None, "punctuation alone quotes nothing"


# -- tier 1 -------------------------------------------------------------------

def scanned(tmp_path, text):
    return claims.scan_claims(doc(tmp_path, text))


def test_stated_numbers_pass_and_invented_ones_fail(tmp_path):
    good, bad = scanned(tmp_path, "ARR is $2.4B, up 42%.\n\nNRR is above 120% since 2023.")
    assert verify.check_numbers(good, facts()) == [] and good.status == "supported"
    assert good.evidence_source == "facts.md" and "$2.4B" in good.quote
    failures = verify.check_numbers(bad, facts())
    assert [(f.raw, f.code) for f in failures] == [("120%", "ONS001"), ("2023", "ONS002")]
    assert bad.status == "unsupported"


def test_arithmetic_on_one_sentence_supports_a_number(tmp_path):
    (c,) = scanned(tmp_path, "Repair turnaround improved 33%.")
    assert verify.check_numbers(c, facts()) == [] and c.status == "derived"
    assert c.derivations == ["33% = 1.4x vs 2.1x in facts.md"], "|1.4 - 2.1| / 2.1"


def test_small_counts_are_their_own_code(tmp_path):
    (c,) = scanned(tmp_path, "Eight of our companies sell to banks.")
    (f,) = verify.check_numbers(c, facts())
    assert f.code == "ONS003"


# -- extraction and merge -----------------------------------------------------

def test_model_spans_must_be_verbatim_and_missed_numbers_are_still_checked(tmp_path):
    d = doc(tmp_path, "ARR reached $2.4B.\n\n"
                      "Customers buy more because onboarding is faster. Churn is 3% and NRR 118%.")
    extracted = [
        claims.ExtractedClaim(block=0, text="ARR reached $2.4B", kind="quantitative", subject="ARR"),
        claims.ExtractedClaim(block=1, text="Customers buy more because onboarding is faster", kind="factual",
                              subject="expansion"),
        claims.ExtractedClaim(block=1, text="Customers love the product", kind="factual", subject="sentiment"),
        claims.ExtractedClaim(block=9, text="NRR 118%", kind="quantitative", subject="NRR"),  # wrong block id
    ]
    out, missed, notes = claims.merge(d, extracted)
    texts = [c.text for c in out]
    assert "Customers love the product" not in texts and "not found verbatim" in notes[0]
    assert "NRR 118%" in texts, "a wrong block id is repaired when the text matches once elsewhere"
    (miss,) = missed
    assert miss.text == "Churn is 3% and NRR 118%." and [q.raw for q in miss.quantities] == ["3%"], \
        "118% sits inside the repaired model claim; 3% is the one the model missed"


def test_one_scan_claim_per_sentence_however_many_numbers(tmp_path):
    d = doc(tmp_path, "Confirm Q3 timing by 15 March.")
    out, missed, _ = claims.merge(d, [])
    assert len(out) == len(missed) == 1 and [q.raw for q in out[0].quantities] == ["Q3", "15 March"]


def test_the_extractor_reads_numbered_blocks(tmp_path):
    seen = {}

    def answer(messages, info: AgentInfo):
        seen["prompt"] = messages[0].parts[-1].content
        return ModelResponse(parts=[ToolCallPart(info.output_tools[0].name, {"claims": [
            {"block": 0, "text": "ARR reached $2.4B", "kind": "quantitative", "subject": "ARR"}]})])

    out = claims.model_claims(doc(tmp_path, "ARR reached $2.4B."), FunctionModel(answer))
    assert out[0].text == "ARR reached $2.4B" and "[0] (line 1) ARR reached $2.4B." in seen["prompt"]


# -- tier 2 -------------------------------------------------------------------

def verdict(status, quote=None, confidence=0.9):
    return {"status": status, "quote": quote, "reason": "because", "confidence": confidence}


def test_a_verdict_without_a_real_quote_is_discarded(tmp_path):
    d = doc(tmp_path, "Revenue doubled.")
    c = claims.Claim(0, d.blocks[0], "Revenue doubled", "factual", "revenue", "model")
    v = verify.apply(c, verify.Verdict(**verdict("supported", "revenue doubled last year")), facts())
    assert v.status == "unsupported" and "not in the evidence" in v.reason
    v = verify.apply(c, verify.Verdict(**verdict("contradicted", "median repair backlog 1.4x")), facts())
    assert v.status == "contradicted" and c.evidence_source == "facts.md"


def test_the_judge_may_clear_a_small_count_but_never_a_number(tmp_path):
    three, invented = scanned(tmp_path, "We hold three plans.\n\nNRR is above 120%.")
    pairs = [(c, verify.check_numbers(c, facts())) for c in (three, invented)]
    assert verify._judgeable(*pairs[0]) and not verify._judgeable(*pairs[1])
    listed = "Plans: east bank, night shuttle, and cargo bikes"
    model = TestModel(custom_output_args=verdict("supported", listed))
    ((claim, v),) = verify.judge(pairs, facts(), model)
    assert claim is three and v.status == "supported"


def test_the_full_pipeline_with_a_scripted_model(tmp_path):
    (tmp_path / "memo.md").write_text(
        "ARR reached $2.4B.\n\nEvery repair hour clears more than an hour of damage.\n\nWe hold three plans.\n\n"
        "Active as of Q4 2025.")

    def answer(messages, info: AgentInfo):
        tool = info.output_tools[0].name
        prompt = messages[0].parts[-1].content
        if tool.endswith("Extraction") or "Document:" in prompt:
            return ModelResponse(parts=[ToolCallPart(tool, {"claims": [
                {"block": 1, "text": "Every repair hour clears more than an hour of damage", "kind": "factual",
                 "subject": "repair capacity"},
                {"block": 2, "text": "We hold three plans", "kind": "quantitative", "subject": "plans"},
            ]})])
        if "three plans" in prompt:
            return ModelResponse(parts=[ToolCallPart(tool, verdict(
                "supported", "Plans: east bank, night shuttle, and cargo bikes"))])
        return ModelResponse(parts=[ToolCallPart(tool, verdict("contradicted", "median repair backlog 1.4x", 0.95))])

    result = check(tmp_path / "memo.md", facts(), Settings(), model=FunctionModel(answer))
    codes = sorted((f.code, f.label) for f in result.findings)
    assert codes == [("ONS002", "line 7"), ("ONS102", "line 3")]
    assert result.used_model


def test_word_documents_read_paragraphs_and_tables_and_serve_as_evidence(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("ARR reached $2.4B.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "NRR", "118%"
    d.save(tmp_path / "memo.docx")
    loaded = documents.load(tmp_path / "memo.docx")
    assert [(b.label, b.text) for b in loaded.blocks] == [
        ("paragraph 1", "ARR reached $2.4B."), ("table 1", "NRR"), ("table 1", "118%")]
    e = ev.load([tmp_path / "memo.docx"])
    assert "118%" in e.passages[0].text and e.passages[0].source == "memo.docx"
